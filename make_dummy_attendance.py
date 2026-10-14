# make_dummy_attendance.py — one-off generator, run with python3.11
"""Create dummy attendance.xlsx for 25 students (24108A0004 onwards) + demo submissions."""
import random
from pathlib import Path
import pandas as pd

BASE = Path(__file__).resolve().parent
DATA = BASE / "data"
DATA.mkdir(exist_ok=True)

# 25 rollnos: 24108A0004 .. 24108A0028
ROLLS = [f"24108A{i:04d}" for i in range(4, 29)]
assert len(ROLLS) == 25, len(ROLLS)

random.seed(42)
rows = []
for r in ROLLS:
    row = {"rollno": r}
    for lab in range(1, 12):
        # ~85% present, deterministic
        row[f"lab{lab}"] = "A" if random.random() < 0.15 else "P"
    rows.append(row)

df = pd.DataFrame(rows, columns=["rollno"] + [f"lab{i}" for i in range(1, 12)])
out = DATA / "attendance.xlsx"
df.to_excel(out, index=False)
print(f"Wrote {out} ({len(df)} students)")
print(df.head(8).to_string(index=False))

# ---- demo submissions with JUMBLED filenames, rollno+assign inside content ----
SW = DATA / "student_work"
SW.mkdir(exist_ok=True)

demos = [
    # (filename, rollno_inside, assign_inside)
    ("random_notes_01.txt", "24108A0004", 1),      # content has both
    ("assign2_24108A0005_report.txt", "24108A0005", 2),  # assign first, then roll
    ("24108A0006_onlyroll.txt", "24108A0006", 3),  # filename only roll, assign in content
    ("mystery_file_xyz.txt", "24108A0007", 1),     # random name, both in content
    ("Lab-4_24108A0008.txt", "24108A0008", 4),
    ("submission_24108A0009_ASSIGNMENT_5.txt", "24108A0009", 5),
    ("24108A0010_lab6_mixed.txt", "24108A0010", 6),
    ("final_doc_student.txt", "24108A0011", 2),    # tests duplicate-assign different student
]

for fname, roll, assign in demos:
    p = SW / fname
    p.write_text(
        f"Roll No: {roll}\nAssignment No: {assign}\n\n"
        f"This is the solution for assignment {assign} by {roll}. "
        f"It explains the experiment, shows code and output.\n",
        encoding="utf-8",
    )
print(f"Wrote {len(demos)} demo submissions into {SW}")

# one .docx demo (needs python-docx) to prove docx path
try:
    from docx import Document
    doc = Document()
    doc.add_paragraph("Roll No: 24108A0012")
    doc.add_paragraph("Lab 3")
    doc.add_paragraph("Solution for lab 3 by 24108A0012 with explanation and output.")
    doc.save(SW / "random_docx_24108A0012.docx")
    print("Wrote random_docx_24108A0012.docx")
except Exception as e:
    print(f"[SKIP] docx demo: {e}")
