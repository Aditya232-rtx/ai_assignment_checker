# marks_mapper.py
"""Map submissions to (rollno x Assign1..11) marks Excel, joined with attendance.

Pipeline:
  1. Scan ONEDRIVE_SYNC_PATH for submission files (.pdf/.docx/.txt/.md/.py/...).
  2. For each file, read text (robust: parser.py first, then python-docx / pymupdf
     fallback since poppler/docx2txt CLIs are often missing) and extract
     (rollno, assignment_no) from CONTENT first, filename as fallback.
     This handles jumbled names: "assign2_24108A0005", "24108A0006_onlyroll",
     "random_notes_01" (both inside doc), etc.
  3. Load attendance.xlsx (rollno | lab1..lab11 with P/A).
  4. Apply rule per cell:
       not submitted            -> 0
       submitted + Present (P)  -> AI score scaled to /10 (or 10 in --no-grade)
       submitted + Absent (A)   -> min(scaled, 5) (or flat 5 in --no-grade)
  5. Write marks_output.xlsx (rollno | Assign1..Assign11 | Total) + print stats.
"""
from __future__ import annotations

import re
from collections import Counter
from pathlib import Path
from typing import Callable, Dict, List, Optional, Tuple

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".py", ".java", ".js", ".cpp", ".c", ".go"}

# Rollno like 24108A0004. Strict pattern + generic fallback + labeled form.
# NOTE: no \b — filenames glue rollnos to underscores (assign2_24108A0005),
# and "_" is a word char so \b fails there. Use digit-aware lookarounds instead.
ROLL_STRICT_RE = re.compile(r"(?<![0-9])\d{5}[A-Z]\d{4}(?![0-9])", re.IGNORECASE)
ROLL_GENERIC_RE = re.compile(r"(?<![0-9])\d{2,5}[A-Z]\d{3,4}(?![0-9])", re.IGNORECASE)
ROLL_LABELED_RE = re.compile(
    r"roll\s*(?:no\.?|number|_?no)?\s*[:\-_\s]*([A-Z0-9]{5,15})", re.IGNORECASE
)

# Assignment/lab number patterns, tried in order. All capture the number.
ASSIGN_RES = [
    re.compile(r"(?:assignment|assign|assgn|assmt|\bass\b)\s*[-_#:.]?\s*(\d{1,2})", re.IGNORECASE),
    re.compile(r"\b(?:lab|expt|experiment|practical|pract)\s*[-_#:.]?\s*(\d{1,2})", re.IGNORECASE),
    re.compile(
        r"(?:exp(?:t)?\.?\s*no\.?|assignment\s*no\.?|assign\s*no\.?|ass\s*no\.?|lab\s*no\.?)\s*[:\-]?\s*(\d{1,2})",
        re.IGNORECASE,
    ),
    # Lab-report headers: "Experiment Number 2", "Lab No 3", "Exp #4".
    # Middle word optional so plain "lab 4"/"assignment 5" still match, but the
    # keyword->number gap is only whitespace/punct so "lab results 5" won't match.
    re.compile(
        r"(?:experiment|expt|exp|lab|assignment|assign|ass)\s*(?:no\.?|number|#)?\s*[:\-]?\s*(\d{1,2})",
        re.IGNORECASE,
    ),
]

# Terse filename-only fallbacks for names like L3_..pdf, A5_..docx, Exp6_...,
# where no full keyword exists. Filename-only on purpose: single letters are
# too ambiguous inside essay/code content (e.g. "L1 cache", "vitamin A").
# The (?<![A-Z0-9]) guard stops matches inside rollnos (24108A0005) and
# words (assign2_...), and v2-style versions can't match (letter blocks it).
SHORT_FILENAME_RE = re.compile(
    r"(?<![A-Z0-9])(?:L|A|EXP?)\s*[-_.#:]?\s*(\d{1,2})(?![0-9])", re.IGNORECASE
)
# Last resort: a bare standalone number in the filename once the rollno is
# stripped out, e.g. 24108A0021_3.pdf. Guarded so years (2024) and v2 can't match.
BARE_NUM_RE = re.compile(r"(?<![A-Z0-9])(\d{1,2})(?![0-9])", re.IGNORECASE)


def _norm_roll(s: str) -> str:
    return re.sub(r"[^A-Z0-9]", "", s.upper())


def extract_rollno(text: str, filename: str = "") -> Optional[str]:
    """Content first, filename fallback. Returns NORMALISED rollno or None."""
    for source in (text, filename):
        if not source:
            continue
        m = ROLL_STRICT_RE.search(source)
        if m:
            return _norm_roll(m.group(0))
        m = ROLL_LABELED_RE.search(source)
        if m:
            cand = _norm_roll(m.group(1))
            # labeled match may grab junk; prefer if it looks like a rollno
            if ROLL_GENERIC_RE.search(cand):
                return cand
        m = ROLL_GENERIC_RE.search(source)
        if m:
            return _norm_roll(m.group(0))
    return None


def extract_assign_no(text: str, filename: str = "", n_assignments: int = 11) -> Optional[int]:
    """Most-common assignment number: content full-keywords -> filename
    full-keywords -> filename short forms (L3/A5/Exp6) -> bare number in
    filename. None if absent/out of range."""
    cands: List[int] = []
    for rx in ASSIGN_RES:
        for m in rx.finditer(text or ""):
            try:
                cands.append(int(m.group(1)))
            except ValueError:
                pass
    if not cands:  # fallback to filename (full keywords first)
        for rx in ASSIGN_RES:
            for m in rx.finditer(filename or ""):
                try:
                    cands.append(int(m.group(1)))
                except ValueError:
                    pass
    if not cands and filename:
        # terse forms: L3, A5, Ex4, Exp6 (filename only — too noisy for content)
        for m in SHORT_FILENAME_RE.finditer(filename):
            try:
                cands.append(int(m.group(1)))
            except ValueError:
                pass
    if not cands and filename:
        # last resort: bare standalone number with the rollno stripped out,
        # e.g. 24108A0021_3.pdf -> 3. Years (2024) and v2 can't match.
        stripped = filename
        for rx in (ROLL_STRICT_RE, ROLL_GENERIC_RE):
            stripped = rx.sub(" ", stripped)
        for m in BARE_NUM_RE.finditer(stripped):
            try:
                cands.append(int(m.group(1)))
            except ValueError:
                pass
    if not cands:
        return None
    num, _ = Counter(cands).most_common(1)[0]
    if 1 <= num <= n_assignments:
        return num
    return None


def read_submission_text(path: Path) -> str:
    """Robust text extraction: parser.py -> python-docx / pymupdf -> raw read."""
    try:
        from .parser import parse_file

        text, _ = parse_file(path)
        if text and text.strip():
            return text
    except Exception:
        pass
    ext = path.suffix.lower()
    try:
        if ext == ".docx":
            from docx import Document

            doc = Document(str(path))
            # Paragraphs PLUS tables (lab reports keep rollno/exp-no in tables).
            # Newline-joined so "Experiment Number\n2" still matches number patterns.
            parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
                    if cells:
                        parts.append(" ".join(cells))
            return "\n".join(parts)
        if ext == ".pdf":
            import fitz  # pymupdf

            with fitz.open(str(path)) as d:
                return "\n".join(page.get_text() for page in d)
    except Exception:
        pass
    try:
        return path.read_text(encoding="utf-8", errors="ignore")
    except Exception:
        return ""


def scan_submissions(
    sync_path: Path,
    n_assignments: int = 11,
    subject: Optional[str] = None,
    exclude: Optional[List[Path]] = None,
) -> Tuple[Dict[Tuple[str, int], Path], List[str]]:
    """Walk sync_path -> {(rollno, assign_no): file_path}, plus unmapped filenames.
    If subject is given, only files with it in the filename (case-insensitive)
    are considered — e.g. subject='mca' picks up MCA_EXP-2.docx, Mca_lab3_AJ_14.docx.
    `exclude` = our own output files (marks_output.txt contains rollnos and
    would otherwise be graded as a submission of its own)."""
    exclude_resolved = set()
    for p in exclude or []:
        try:
            exclude_resolved.add(Path(p).resolve())
        except OSError:
            exclude_resolved.add(Path(p))
    found: Dict[Tuple[str, int], Path] = {}
    unmapped: List[str] = []
    dupes: List[str] = []
    skipped = 0
    for fp in sorted(sync_path.rglob("*")):
        if not fp.is_file() or fp.suffix.lower() not in ALLOWED_EXTENSIONS:
            continue
        try:
            if fp.resolve() in exclude_resolved:
                continue
        except OSError:
            continue
        if subject and subject.lower() not in fp.name.lower():
            skipped += 1
            continue
        text = read_submission_text(fp)
        roll = extract_rollno(text, fp.name)
        ano = extract_assign_no(text, fp.name, n_assignments)
        if roll and ano:
            key = (roll, ano)
            if key in found:
                dupes.append(f"{fp.name} duplicates {found[key].name} for {roll}/A{ano}")
            else:
                found[key] = fp
        else:
            why = []
            if not roll:
                why.append("no rollno in content/filename")
            if not ano:
                why.append("no lab/assign no. in content/filename")
            if not (text or "").strip():
                why.append("empty/unreadable text (scanned PDF?)")
            unmapped.append(f"{fp.name} (roll={roll}, assign={ano}; {'; '.join(why)})")
    for d in dupes:
        print(f"[DUPE] {d}")
    if subject:
        print(f"[FILTER] subject='{subject}': {skipped} non-matching files skipped")
    return found, unmapped


def find_attendance(sync_path: Path) -> Optional[Path]:
    """Auto-locate the attendance workbook inside the subject folder.

    Prefers a top-level *attendance*.xls* file, else the first recursive match.
    Returns None when there is no attendance sheet yet.
    """
    top = sorted(
        p for p in sync_path.iterdir()
        if p.is_file() and "attendance" in p.name.lower() and p.suffix.lower() in (".xlsx", ".xls")
    )
    if top:
        return top[0]
    deep = sorted(
        p for p in sync_path.rglob("*")
        if p.is_file() and "attendance" in p.name.lower() and p.suffix.lower() in (".xlsx", ".xls")
    )
    return deep[0] if deep else None


def _norm_attendance_val(v) -> str:
    """Decode attendance cells. Supports P/A words AND numeric codes:
    professor sheets use 0 = absent, 2 = present (also 1 = present)."""
    s = str(v).strip().upper()
    if s in ("P", "PRESENT", "1", "2", "YES", "Y"):
        return "P"
    if s in ("A", "ABSENT", "0", "NO", "N"):
        return "A"
    # Floats from Excel: "2.0" -> P, "0.0" -> A
    try:
        num = float(s)
        if num >= 1:
            return "P"
        return "A"
    except ValueError:
        pass
    return "P" if s.startswith("P") else ("A" if s.startswith("A") else "P")


def _roll_col(df) -> str:
    """Find the rollno column by fuzzy name; else fall back to first column."""
    aliases = ("rollno", "roll no", "roll no.", "roll_no", "roll", "roll number", "id", "student id", "studentid")
    for c in df.columns:
        if c in aliases:
            return c
    for c in df.columns:
        if "roll" in c or "student id" in c:
            return c
    return df.columns[0]


def _lab_col_map(df, n_assignments: int) -> Dict[int, str]:
    """Map lab number 1..N -> actual column name. Fuzzy: lab1, lab 1, exp2,
    expt 3, ex-1 (professor's header), experiment4, assign5, l6, bare '7'."""
    mapping: Dict[int, str] = {}
    for c in df.columns:
        cleaned = re.sub(r"[^a-z0-9]", "", c)
        m = re.match(
            r"^(?:lab|ex|exp|expt|experiment|assign|assignment|ass|l)(\d{1,2})$",
            cleaned,
        )
        if not m and re.fullmatch(r"\d{1,2}", cleaned):
            m = re.match(r"(\d{1,2})", cleaned)  # bare numeric header
        if m:
            num = int(m.group(1))
            if 1 <= num <= n_assignments and num not in mapping:
                mapping[num] = c
    return mapping


def load_attendance(attendance_xlsx: Path, n_assignments: int = 11):
    """Returns (rollnos_in_order, {roll: {lab_no: 'P'/'A'}}).
    Tolerates prof-sheet quirks: 0/2 numeric codes, column-name variants."""
    import pandas as pd

    df = pd.read_excel(attendance_xlsx, dtype=str)
    df.columns = [str(c).strip().lower() for c in df.columns]
    roll_col = _roll_col(df)
    lab_cols = _lab_col_map(df, n_assignments)
    att: Dict[str, Dict[int, str]] = {}
    order: List[str] = []
    for _, row in df.iterrows():
        roll = _norm_roll(str(row[roll_col]))
        if not roll or roll == "NAN":
            continue
        order.append(roll)
        att[roll] = {}
        for lab in range(1, n_assignments + 1):
            if lab in lab_cols:
                att[roll][lab] = _norm_attendance_val(row[lab_cols[lab]])
            else:
                att[roll][lab] = "P"  # missing column -> assume present
    return order, att


# Hidden worksheet inside marks_output.xlsx that stores per-cell AI grades
# from previous runs. Keyed on file sha256 + grading meta so a re-run only
# re-grades NEW or CHANGED submissions — and the same two output files
# (marks_output.xlsx / marks_output.txt) are updated in place every time.
CACHE_SHEET = "AI_Grades"


def _file_sha256(path: Path) -> str:
    import hashlib

    h = hashlib.sha256()
    try:
        with path.open("rb") as f:
            for chunk in iter(lambda: f.read(65536), b""):
                h.update(chunk)
    except OSError:
        return ""
    return h.hexdigest()


def rubric_signature(rubric: dict) -> str:
    """Short fingerprint of the rubric criteria — part of the cache meta so
    editing the rubric triggers re-grading instead of reusing stale scores."""
    import hashlib
    import json

    blob = json.dumps(rubric.get("criteria", []), sort_keys=True, default=str)
    return hashlib.sha256(blob.encode("utf-8")).hexdigest()[:12]


def load_grade_cache(output_xlsx: Path) -> Dict[Tuple[str, int], Dict[str, str]]:
    """Read previously computed raw AI grades from the hidden AI_Grades sheet
    of the existing output workbook. Missing/locked/unreadable workbook ->
    empty cache (safe fallback: everything gets re-graded)."""
    import pandas as pd

    if not output_xlsx.is_file():
        return {}
    try:
        df = pd.read_excel(output_xlsx, sheet_name=CACHE_SHEET, dtype=str)
    except FileNotFoundError:
        return {}
    except ValueError:
        # workbook exists but has no AI_Grades sheet yet (first run / old file)
        return {}
    except Exception as e:
        print(f"[CACHE] previous grades not readable ({type(e).__name__}); re-grading all")
        return {}
    cache: Dict[Tuple[str, int], Dict[str, str]] = {}
    for _, r in df.iterrows():
        try:
            roll = str(r.get("rollno", "")).strip().lower()
            if not roll or roll in ("nan", "none", ""):
                continue
            key = (str(r["rollno"]).strip(), int(float(str(r["assign"]))))
            cache[key] = {
                "raw": str(r.get("raw_score", "")),
                "why": str(r.get("reasoning", "")),
                "sha256": str(r.get("sha256", "")),
                "meta": str(r.get("meta", "")),
                "file": str(r.get("file", "")),
            }
        except Exception:
            continue
    return cache


def default_score_fn(_path: Path, _text: str, _assign_no: int) -> Tuple[float, str]:
    """No-LLM scoring: full 10 (attendance rule caps absent to 5 later)."""
    return 10.0, "Deterministic mode (--no-grade): submitted work counted as 10/10 before the attendance rule."


def _criterion_reason(grading: List[Dict[str, Any]]) -> str:
    """Flatten per-criterion AI justifications into one line for the report."""
    parts = []
    for g in grading:
        just = str(g.get("justification", "")).replace("\n", " ").strip()
        ev = str(g.get("evidence", "")).replace("\n", " ").strip()
        parts.append(
            f"{g.get('criterion','?')}: {g.get('score','?')}/{g.get('max','?')} — {just} (evidence: {ev})"
        )
    return " | ".join(parts) if parts else "AI returned no criteria."


def ai_score_fn_factory(rubric: dict) -> Callable[[Path, str, int], Tuple[float, str]]:
    """Score 0..10 via grader agent (scaled total/max*10). Falls back to 10.
    Returns (score, reasoning) so the marks report can explain each mark."""

    def _fn(path: Path, text: str, _assign_no: int) -> Tuple[float, str]:
        try:
            from .agents.grader import grade_submission

            grading = grade_submission(rubric, text[:2000])
            total = sum(g.get("score", 0) for g in grading)
            mx = sum(g.get("max", 0) for g in grading) or 1
            score = round(max(0.0, min(1.0, total / mx)) * 10.0, 2)
            reason = _criterion_reason(grading)
            return score, f"AI rubric grading: {reason} → raw {score:.2f}/10"
        except Exception as e:
            print(f"[WARN] AI grade failed for {path.name}: {e} -> 10")
            return 10.0, f"AI grading failed ({e}); defaulted to 10/10 before the attendance rule."

    return _fn


def build_marks(
    attendance_xlsx: Path,
    sync_path: Path,
    output_xlsx: Path,
    n_assignments: int = 11,
    score_fn: Optional[Callable[..., Any]] = None,
    subject: Optional[str] = None,
    write_reasoning: bool = True,
    grade_meta: Optional[str] = None,
):
    """Full mapper: scan -> join -> apply 10/5/0 rule -> write Excel (+ a
    .txt reasoning report next to it). Returns the marks DataFrame.

    Incremental: previously graded cells (same file sha256 + same grade_meta)
    are read back from the hidden AI_Grades sheet and reused — only new or
    changed submissions hit the AI. Both output files are overwritten in
    place; no new files are created on re-runs."""
    import pandas as pd
    from datetime import datetime

    from .config import MODEL_NAME

    score_fn = score_fn or default_score_fn
    if grade_meta is None:
        grade_meta = "nograde" if score_fn is default_score_fn else f"model={MODEL_NAME}"
    cache = load_grade_cache(output_xlsx)
    reused = 0
    graded = 0
    ai_rows: List[Dict[str, Any]] = []  # full cache sheet contents (old + new)
    order, att = load_attendance(attendance_xlsx, n_assignments)
    submissions, unmapped = scan_submissions(
        sync_path,
        n_assignments,
        subject,
        exclude=[output_xlsx, output_xlsx.with_suffix(".txt")],
    )

    print(f"[SCAN] {len(submissions)} mapped submissions, {len(unmapped)} unmapped")
    for u in unmapped:
        print(f"  [UNMAPPED] {u}")

    # all rolls = attendance rolls + submission-only rolls (late/new admissions)
    all_rolls: List[str] = list(order)
    for (roll, _a) in submissions:
        if roll not in att:
            att[roll] = {lab: "P" for lab in range(1, n_assignments + 1)}
            all_rolls.append(roll)
            print(f"[NEW] {roll} submitted but not in attendance -> treated as Present")

    rows = []
    entries: List[str] = []   # per-cell reasoning blocks for the .txt report
    zeros: List[str] = []     # missing-submission notes
    for roll in all_rolls:
        row = {"rollno": roll}
        missing_labs: List[int] = []
        for lab in range(1, n_assignments + 1):
            key = (roll, lab)
            if key not in submissions:
                row[f"Assign{lab}"] = 0
                missing_labs.append(lab)
                continue
            present = att[roll].get(lab, "P") == "P"
            fp = submissions[key]
            sha = _file_sha256(fp)
            cached = cache.get(key) if sha else None
            if cached and cached["sha256"] == sha and cached["meta"] == grade_meta:
                # unchanged file + same rubric/model -> reuse, skip AI call
                raw10 = max(0.0, min(10.0, float(cached["raw"])))
                why = cached["why"]
                display_why = f"{why} [reused from previous run: file unchanged]"
                reused += 1
            else:
                text = read_submission_text(fp)
                res = score_fn(fp, text, lab)
                if isinstance(res, tuple):
                    raw10, why = float(res[0]), str(res[1])
                else:
                    raw10, why = float(res), "No reasoning returned by scorer."
                raw10 = max(0.0, min(10.0, raw10))
                display_why = why
                graded += 1
            final = raw10 if present else min(raw10, 5.0)
            row[f"Assign{lab}"] = final
            ai_rows.append(
                {
                    "rollno": roll,
                    "assign": lab,
                    "raw_score": raw10,
                    "reasoning": why[:32000],
                    "file": fp.name,
                    "sha256": sha,
                    "meta": grade_meta,
                    "graded_at": datetime.now().isoformat(timespec="seconds"),
                }
            )

            if present:
                rule = f"Present + submitted -> mark = AI raw score = {final:.2f}/10"
            elif raw10 > 5.0:
                rule = (
                    f"Absent + submitted -> rule caps mark at 5: "
                    f"min(AI raw {raw10:.2f}, 5) = {final:.2f}/10"
                )
            else:
                rule = (
                    f"Absent + submitted -> rule allows at most 5: "
                    f"min(AI raw {raw10:.2f}, 5) = {final:.2f}/10 (raw already within cap)"
                )
            entries.append(
                f"[{roll} | Expt {lab}]\n"
                f"  File:       {fp.name}\n"
                f"  Attendance: {'Present' if present else 'Absent'}\n"
                f"  Grading:    {display_why}\n"
                f"  Rule:       {rule}\n"
                f"  FINAL MARK: {final:.2f} / 10\n"
            )
        if missing_labs:
            labs_str = ", ".join(f"Expt {l}" for l in missing_labs)
            zeros.append(f"  {roll}: no submission file found for {labs_str} -> 0/10")
        rows.append(row)

    cols = ["rollno"] + [f"Assign{i}" for i in range(1, n_assignments + 1)]
    df = pd.DataFrame(rows, columns=cols)
    df["Total"] = df[[c for c in cols if c != "rollno"]].sum(axis=1).round(2)

    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_xlsx, engine="openpyxl") as w:
        df.to_excel(w, index=False, sheet_name="Marks")
        ws = w.sheets["Marks"]
        ws.freeze_panes = "B2"
        widths = [14] + [9] * n_assignments + [10]
        for i, wd in enumerate(widths, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = wd
        # hidden cache sheet: raw AI grades for incremental re-runs
        cache_df = pd.DataFrame(
            ai_rows,
            columns=[
                "rollno", "assign", "raw_score", "reasoning",
                "file", "sha256", "meta", "graded_at",
            ],
        )
        cache_df.to_excel(w, index=False, sheet_name=CACHE_SHEET)
        w.sheets[CACHE_SHEET].sheet_state = "hidden"

    if write_reasoning:
        report_path = output_xlsx.with_suffix(".txt")
        lines = [
            "MARKS REASONING REPORT",
            "=" * 60,
            f"Generated:  {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
            f"Folder:     {sync_path}",
            f"Attendance: {attendance_xlsx.name}  (P/2 = present, A/0 = absent)",
            f"Grader:     AI rubric grading via Ollama ({MODEL_NAME})",
            f"Run stats:  {reused} cells reused from previous runs (file unchanged), "
            f"{graded} newly graded",
            "",
            "NOTE: this report and marks_output.xlsx are updated IN PLACE on every",
            "run — new/changed submissions are picked up, unchanged ones are reused.",
            "",
            "MARKING RULE:",
            "  * Not submitted                 -> 0/10",
            "  * Present (P/2) + submitted     -> AI score out of 10",
            "  * Absent  (A/0) + submitted     -> min(AI score, 5)",
            "",
            "-" * 60,
            "PER-CELL REASONING",
            "-" * 60,
            "",
        ]
        lines.extend(entries)
        if zeros:
            lines += [
                "-" * 60,
                "NO SUBMISSION FOUND (0/10)",
                "-" * 60,
                *zeros,
                "",
            ]
        if unmapped:
            lines += [
                "-" * 60,
                "SKIPPED FILES (rollno/expt no. could not be read)",
                "-" * 60,
                *[f"  {u}" for u in unmapped],
                "",
            ]
        report_path.write_text("\n".join(lines), encoding="utf-8")
        print(f"[DONE] reasoning report: {report_path}")

    submitted_cells = sum(
        1 for r in rows for lab in range(1, n_assignments + 1) if r[f"Assign{lab}"] > 0
    )
    print(
        f"[DONE] {output_xlsx} : {len(df)} students x {n_assignments} assigns, "
        f"{submitted_cells} non-zero cells | [CACHE] reused {reused}, newly graded {graded}"
    )
    return df
